"""DOCX 보고서 빌더 핵심 패턴 — v3 부암동 보고서에서 추출.

이 파일은 새 매물 보고서 생성 시 복사해서 사용할 수 있는
실제 작동 코드 패턴입니다. {{변수}} 부분만 교체하면 됩니다.
"""

import os, re, sys, json
from pathlib import Path

_USER_SP = "/home/grafy/.local/lib/python3.12/site-packages"
if Path(_USER_SP).exists() and _USER_SP not in sys.path:
    sys.path.insert(0, _USER_SP)

from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap

# ── 스타일 상수 ──
FONT_KR = "Malgun Gothic"
PRIMARY = RGBColor(0x1F, 0x29, 0x4B)       # 진한 남색
ACCENT = RGBColor(0xC0, 0x39, 0x2B)        # 강조 적색
GRAY = RGBColor(0x58, 0x58, 0x58)
LIGHT_GRAY_HEX = "E8ECF3"
BG_CALLOUT_HEX = "F2F4F8"
WARN_HEX = "FFF4E6"
SUCCESS_HEX = "E9F6EC"
PAGE_WIDTH_IN = 6.0   # 단독 이미지 폭 (2열=3.0, 4그리드=2.8)
CONCLUSION_GREEN = RGBColor(0x2D, 0x4A, 0x3E)

# ── 약어 치환 (glossary.md §3) ──
DIMENSION_REPLACEMENTS = {
    'LOC': '입지·접근성', 'PRC': '가격 적정성', 'PHY': '물리적 상태',
    'REG': '법적·규제', 'MKT': '상권·문화', 'FIT': '목적 적합도',
    'SCR': '희소성·미래가치', 'FIN': '재무·수익성',
}
EXTRA_REPLACEMENTS = {
    'CMA': '비교 시장 분석', 'ROI': '투자 수익률',
    'NOI': '순영업수익', 'GB': '개발제한구역',
}

def expand_abbreviations(text: str) -> str:
    """단어 경계를 지켜 치환. LOCAL·LOCK 같은 건 제외."""
    if not isinstance(text, str):
        return text
    for abbr, full in {**DIMENSION_REPLACEMENTS, **EXTRA_REPLACEMENTS}.items():
        text = re.sub(rf'\b{abbr}\b', full, text)
    return text


# ── 저수준 XML 헬퍼 ──
def _set_cell_shading(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)

def _set_cell_border(cell, color="BFBFBF", size="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), size)
        b.set(qn("w:color"), color)
        tcBorders.append(b)
    tcPr.append(tcBorders)

def _run_font(run, size=None, bold=None, color=None, font_name=FONT_KR, italic=None):
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:cs"), font_name)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color is not None: run.font.color.rgb = color
    if italic is not None: run.italic = italic


# ── 하이퍼링크 (python-docx 미지원이라 XML 직접 생성) ──
def add_hyperlink(paragraph, url, text, color="0563C1", underline=True, size=10, bold=False):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), FONT_KR)
    rFonts.set(qn("w:hAnsi"), FONT_KR)
    rFonts.set(qn("w:eastAsia"), FONT_KR)
    rPr.append(rFonts)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(size * 2)); rPr.append(sz)
    szCs = OxmlElement("w:szCs"); szCs.set(qn("w:val"), str(size * 2)); rPr.append(szCs)
    c = OxmlElement("w:color"); c.set(qn("w:val"), color); rPr.append(c)
    if underline:
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    if bold:
        b = OxmlElement("w:b"); rPr.append(b)
    new_run.append(rPr)
    t = OxmlElement("w:t"); t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


# ── 문서 초기화 ──
def init_document():
    doc = Document()
    for section in doc.sections:
        section.page_height = Cm(29.7)
        section.page_width = Cm(21.0)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
    style = doc.styles["Normal"]
    style.font.name = FONT_KR
    style.font.size = Pt(10)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), FONT_KR)
    rFonts.set(qn("w:hAnsi"), FONT_KR)
    rFonts.set(qn("w:eastAsia"), FONT_KR)
    rFonts.set(qn("w:cs"), FONT_KR)
    return doc


# ── 고수준 빌딩 블록 ──
def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)

def add_para(doc, text="", size=10, bold=False, color=None, align=None, italic=False, space_after=4):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    if space_after is not None: p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        _run_font(r, size=size, bold=bold, color=color, italic=italic)
    return p

def add_heading(doc, text, level=1):
    """Word Heading 스타일 적용 + 시각 속성 덮어쓰기."""
    sizes = {1: 16, 2: 13, 3: 11}
    p = doc.add_heading("", level=level)
    p.paragraph_format.space_before = Pt(6 if level > 1 else 12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    _run_font(r, size=sizes.get(level, 10), bold=True, color=PRIMARY)
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "1F294B")
        pBdr.append(bottom)
        pPr.append(pBdr)
    return p

def add_conclusion_box(doc, text):
    """✅ 결론 박스: 딥그린 배경 + 흰색 bold — 12섹션 각각 끝에 필수."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    _set_cell_shading(cell, "2D4A3E")
    _set_cell_border(cell, color="2D4A3E", size="8")
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(2)
    content = text if text.startswith("✅") else f"✅ 결론: {text}"
    r = para.add_run(content)
    _run_font(r, size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return tbl

def add_summary_box(doc, text, fill_hex=BG_CALLOUT_HEX, bold=True, color=PRIMARY, size=10):
    """핵심요약 박스: 1x1 회색 배경 테이블."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    _set_cell_shading(cell, fill_hex)
    _set_cell_border(cell, color="B8C0D0", size="6")
    cell.text = ""
    para = cell.paragraphs[0]
    for i, line in enumerate(text.split("\n")):
        if i > 0: para = cell.add_paragraph()
        r = para.add_run(line)
        _run_font(r, size=size, bold=bold, color=color)
        para.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl

def add_table(doc, headers, rows, col_widths_in=None, header_fill="1F294B", body_fill=None):
    """표 생성. rows 값에 (text, url) 튜플 허용 (하이퍼링크)."""
    ncols = len(headers)
    tbl = doc.add_table(rows=1 + len(rows), cols=ncols)
    tbl.style = "Light Grid Accent 1"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    hdr = tbl.rows[0].cells
    for i, h in enumerate(headers):
        c = hdr[i]; c.text = ""
        _set_cell_shading(c, header_fill)
        _set_cell_border(c, color="1F294B", size="8")
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(h))
        _run_font(r, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for ri, row in enumerate(rows):
        cells = tbl.rows[ri + 1].cells
        for ci, val in enumerate(row):
            c = cells[ci]; c.text = ""
            _set_cell_border(c, color="BFBFBF", size="4")
            if body_fill and ri % 2 == 1:
                _set_cell_shading(c, body_fill)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            if isinstance(val, tuple):
                if val[1]: add_hyperlink(p, val[1], val[0], size=9)
                else:
                    r = p.add_run(val[0]); _run_font(r, size=9)
            else:
                text = str(val) if val is not None else ""
                if text.startswith("**") and text.endswith("**") and len(text) > 4:
                    r = p.add_run(text[2:-2]); _run_font(r, size=9, bold=True)
                else:
                    r = p.add_run(text); _run_font(r, size=9)
    if col_widths_in:
        _sync_tbl_grid(tbl, col_widths_in)
        for row in tbl.rows:
            for ci, w in enumerate(col_widths_in):
                if ci < len(row.cells): row.cells[ci].width = Inches(w)
    return tbl

def _sync_tbl_grid(tbl, col_widths_in):
    tbl_element = tbl._tbl
    grid = tbl_element.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid"); tbl_element.insert(0, grid)
    for gc in list(grid.findall(qn("w:gridCol"))): grid.remove(gc)
    for w_in in col_widths_in:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(round(w_in * 1440))))
        grid.append(gc)


# ── 이미지 처리 (EXIF 보정 + 중복 검출) ──
USED_IMAGES = []
IMAGE_PHASHES = {}

def _phash(im, size=16):
    """dhash — 이웃 픽셀 차이 기반, 구조적 변화에 민감."""
    from PIL import Image
    g = im.convert("L").resize((size + 1, size), Image.LANCZOS)
    px = list(g.getdata())
    bits = []
    for y in range(size):
        for x in range(size):
            left = px[y * (size + 1) + x]
            right = px[y * (size + 1) + x + 1]
            bits.append("1" if left > right else "0")
    return "".join(bits)

def _optimize_image(src_path, img_cache, max_width=1800, quality=85):
    """EXIF 회전 보정 + 리사이즈 + 캐시."""
    from PIL import Image, ImageOps
    src = Path(src_path)
    if not src.exists(): return src
    img_cache.mkdir(parents=True, exist_ok=True)
    cached = img_cache / src.name
    if cached.exists() and cached.stat().st_mtime >= src.stat().st_mtime:
        if src.name not in USED_IMAGES: USED_IMAGES.append(src.name)
        return cached
    try:
        im = Image.open(src)
        im = ImageOps.exif_transpose(im)
        if im.mode not in ("RGB", "L"): im = im.convert("RGB")
        w, h = im.size
        if w > max_width:
            new_h = int(h * max_width / w)
            im = im.resize((max_width, new_h), Image.LANCZOS)
        IMAGE_PHASHES[src.name] = _phash(im)
        im.save(cached, "JPEG", quality=quality, optimize=True, progressive=True)
        if src.name not in USED_IMAGES: USED_IMAGES.append(src.name)
        return cached
    except Exception as e:
        print(f"[WARN] image optimize failed for {src.name}: {e}")
        return src

def add_image(doc, path, img_cache, caption=None, width_in=PAGE_WIDTH_IN):
    if not Path(path).exists():
        add_para(doc, f"[이미지 누락: {path}]", size=9, color=ACCENT, italic=True)
        return
    if width_in > 6.2: width_in = 6.2
    opt = _optimize_image(path, img_cache)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    r.add_picture(str(opt), width=Inches(width_in))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(4)
        rc = cp.add_run(caption)
        _run_font(rc, size=9, italic=True, color=GRAY)

def add_image_grid(doc, image_paths_captions, img_cache, cols=2, img_width_in=3.15):
    """2xN 이미지 그리드."""
    rows = (len(image_paths_captions) + cols - 1) // cols
    tbl = doc.add_table(rows=rows * 2, cols=cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (path, caption) in enumerate(image_paths_captions):
        r = (idx // cols) * 2
        c = idx % cols
        img_cell = tbl.rows[r].cells[c]
        cap_cell = tbl.rows[r + 1].cells[c]
        img_cell.text = ""
        p = img_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if Path(path).exists():
            opt = _optimize_image(path, img_cache)
            run = p.add_run()
            run.add_picture(str(opt), width=Inches(img_width_in))
        else:
            run = p.add_run(f"[누락: {Path(path).name}]")
            _run_font(run, size=8, color=ACCENT, italic=True)
        cap_cell.text = ""
        cp = cap_cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        txt = caption.strip()
        if not txt.startswith("▲"): txt = "▲ " + txt
        rc = cp.add_run(txt)
        _run_font(rc, size=8, italic=True, color=GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ── Soul 카드 그리드 (2열 × 3행 = 6장) ──
def build_soul_cards_grid(doc, soul_cards, img_dir, img_cache):
    """Soul 6 카드: 이미지 + 제목 + 설명 + 출처 하이퍼링크."""
    cols = 2
    rows = (len(soul_cards) + cols - 1) // cols
    tbl = doc.add_table(rows=rows * 3, cols=cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, card in enumerate(soul_cards):
        r = (idx // cols) * 3
        c = idx % cols
        img_cell = tbl.rows[r].cells[c]
        ttl_cell = tbl.rows[r + 1].cells[c]
        src_cell = tbl.rows[r + 2].cells[c]
        # 이미지
        img_cell.text = ""
        _set_cell_border(img_cell, color="D8DCE5", size="4")
        p = img_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_path = img_dir / card["image_file"]
        if img_path.exists():
            opt = _optimize_image(img_path, img_cache)
            run = p.add_run()
            run.add_picture(str(opt), width=Inches(3.0))
        else:
            run = p.add_run(f"[누락: {card['image_file']}]")
            _run_font(run, size=8, color=ACCENT, italic=True)
        # 제목 + 설명
        ttl_cell.text = ""
        _set_cell_border(ttl_cell, color="D8DCE5", size="4")
        _set_cell_shading(ttl_cell, "F9FAFC")
        tp = ttl_cell.paragraphs[0]
        tp.paragraph_format.space_after = Pt(2)
        r_ = tp.add_run(f"💎 {card['title']}")
        _run_font(r_, size=10, bold=True, color=PRIMARY)
        dp = ttl_cell.add_paragraph()
        dp.paragraph_format.space_after = Pt(2)
        rd = dp.add_run(card["description"])
        _run_font(rd, size=8, color=GRAY)
        # 출처 링크
        src_cell.text = ""
        _set_cell_border(src_cell, color="D8DCE5", size="4")
        _set_cell_shading(src_cell, "F2F4F8")
        sp = src_cell.paragraphs[0]
        sp.paragraph_format.space_after = Pt(1)
        rs = sp.add_run("출처: ")
        _run_font(rs, size=7, bold=True, color=GRAY)
        for i, s in enumerate(card.get("sources", [])):
            url = s.get("url", "")
            label = s.get("label", "")
            if url.startswith("http"):
                if i > 0:
                    spc = sp.add_run(" · ")
                    _run_font(spc, size=7, color=GRAY)
                add_hyperlink(sp, url, label, size=7)
            else:
                if i > 0:
                    spc = sp.add_run(" · ")
                    _run_font(spc, size=7, color=GRAY)
                rlit = sp.add_run(label)
                _run_font(rlit, size=7, italic=True, color=GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ── 용어집 1×8 표 ──
def add_glossary_table(doc):
    """점수 차원 용어집 1x8 표 (첫 점수 페이지 상단)."""
    add_heading(doc, "점수 차원 용어집 — 8개 평가 차원", level=2)
    rows_data = [
        ("입지·접근성", "이동 편의"), ("가격 적정성", "실거래 대비 포지셔닝"),
        ("물리적 상태", "건물·설비 컨디션"), ("법적·규제", "활용 가능성"),
        ("상권·문화", "주변 맥락"), ("목적 적합도", "용도 매칭"),
        ("희소성·미래가치", "공급 제약 · 스토리"), ("재무·수익성", "비용·수익 효율"),
    ]
    tbl = doc.add_table(rows=2, cols=8)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ci, (name, _) in enumerate(rows_data):
        c = tbl.rows[0].cells[ci]; c.text = ""
        _set_cell_shading(c, "1F294B")
        _set_cell_border(c, color="1F294B", size="8")
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(name)
        _run_font(r, size=8, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for ci, (_, desc) in enumerate(rows_data):
        c = tbl.rows[1].cells[ci]; c.text = ""
        _set_cell_shading(c, "F2F4F8")
        _set_cell_border(c, color="B8C0D0", size="4")
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(desc)
        _run_font(r, size=8, color=GRAY)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for row in tbl.rows:
        for ci in range(8): row.cells[ci].width = Inches(0.85)
    _sync_tbl_grid(tbl, [0.85] * 8)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ── main() 구조 ──
def main_structure():
    """실제 main() 호출 순서 — 이 순서대로 섹션을 빌드하면 됩니다."""
    return [
        "build_cover(doc)",
        "page_break(doc)",
        "build_score_header(doc, soul_cards)",   # 점수 카드 4p
        "build_executive_snapshot(doc)",          # 핵심 요약
        "build_value_section(doc)",               # 가치 판단
        "build_priority_checklist(doc)",           # 우선순위 체크리스트
        "build_section_1_basic(doc)",              # §1 매물 개요
        "build_section_2_location(doc)",           # §2 입지·접근성
        "build_section_3_price(doc)",              # §3 가격 적정성
        "build_section_4_market(doc, sec_manifest)",  # §4 상권·문화
        "build_section_5_acquisition_tax(doc)",     # §5 취득세
        "build_section_6_holding_tax(doc)",         # §6 보유세
        "build_section_7_usage_change(doc)",        # §7 용도변경
        "build_section_8_interior(doc, sec_manifest)",  # §8 인테리어
        "build_section_9_rebuild(doc)",             # §9 재건축
        "build_section_10_risk(doc)",               # §10 리스크
        "build_section_11_operations(doc, sec_manifest)",  # §11 운영
        "build_section_12_conclusion(doc)",         # §12 종합
        "build_appendix_a(doc)",                   # A: 데이터 출처
        "build_appendix_b(doc)",                   # B: CMA 비교군
        "build_appendix_c(doc)",                   # C: 인테리어 레퍼런스
        "build_appendix_d(doc)",                   # D: 민감도·신뢰도 (메타→Appendix)
    ]
