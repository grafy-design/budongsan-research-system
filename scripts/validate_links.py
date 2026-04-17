#!/usr/bin/env python3
"""
validate_links.py -- 외부 URL 전수 검증 스크립트

입력: docx 파일 또는 HTML 파일
출력: 검증 리포트 JSON

사용법:
  python validate_links.py --docx output.docx --output report.json
  python validate_links.py --html index.html --output report.json
  python validate_links.py --docx output.docx --html index.html --output report.json
"""
import argparse
import json
import re
import subprocess
import sys
import time
from pathlib import Path
from urllib.parse import urlparse

# ---- DOCX 하이퍼링크 추출 ----
def extract_links_from_docx(docx_path: str) -> list:
    """python-docx로 docx 내 하이퍼링크 추출"""
    try:
        from docx import Document
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
    except ImportError:
        print("오류: python-docx 패키지가 필요합니다. pip install python-docx", file=sys.stderr)
        sys.exit(1)

    doc = Document(docx_path)
    links = []

    # 본문 paragraph 내 하이퍼링크
    for para in doc.paragraphs:
        for run in para.runs:
            # run 자체에서 하이퍼링크 탐색
            pass

        # XML에서 직접 하이퍼링크 추출
        for hyperlink in para._element.findall(
            ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink"
        ):
            # 외부 링크: r:id 속성 참조
            r_id = hyperlink.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            if r_id:
                try:
                    rel = doc.part.rels[r_id]
                    url = rel.target_ref
                    # 링크 텍스트 추출
                    texts = hyperlink.findall(
                        ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
                    )
                    text = "".join(t.text or "" for t in texts)
                    links.append({
                        "url": url,
                        "text": text.strip(),
                        "source": f"docx:{docx_path}",
                        "type": "hyperlink",
                    })
                except KeyError:
                    pass

    # 표 내 하이퍼링크
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for hyperlink in para._element.findall(
                        ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink"
                    ):
                        r_id = hyperlink.get(
                            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                        )
                        if r_id:
                            try:
                                rel = cell.part.rels[r_id]
                                url = rel.target_ref
                                texts = hyperlink.findall(
                                    ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
                                )
                                text = "".join(t.text or "" for t in texts)
                                links.append({
                                    "url": url,
                                    "text": text.strip(),
                                    "source": f"docx:{docx_path}:table",
                                    "type": "hyperlink",
                                })
                            except KeyError:
                                pass

    return links


# ---- HTML 하이퍼링크 추출 ----
def extract_links_from_html(html_path: str) -> list:
    """BeautifulSoup 또는 regex로 HTML 내 href 추출"""
    content = Path(html_path).read_text(encoding="utf-8")
    links = []

    try:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(content, "html.parser")
        for a in soup.find_all("a", href=True):
            url = a["href"]
            if url.startswith(("http://", "https://")):
                links.append({
                    "url": url,
                    "text": a.get_text(strip=True),
                    "source": f"html:{html_path}",
                    "type": "href",
                })
    except ImportError:
        # BeautifulSoup 없으면 regex fallback
        pattern = r'href=["\']?(https?://[^"\'>\s]+)["\']?'
        for match in re.finditer(pattern, content):
            links.append({
                "url": match.group(1),
                "text": "",
                "source": f"html:{html_path}",
                "type": "href_regex",
            })

    return links


# ---- URL 검증 ----
# 매물 링크의 컨텐츠 키워드 (지역 검색 루트로 빠지는지 확인용)
PROPERTY_KEYWORDS = ["면적", "가격", "매매", "전용", "평", "호가", "매물", "보증금", "월세", "전세",
                     "sqm", "m²", "price", "area"]
REDIRECT_KEYWORDS = ["지역 검색", "검색 결과", "매물 목록", "전체 매물", "리스트"]


def check_url(url: str, timeout: int = 15) -> dict:
    """curl로 URL 상태 확인 + 컨텐츠 키워드 매칭"""
    result = {
        "url": url,
        "status_code": None,
        "final_url": None,
        "ok": False,
        "redirect": False,
        "content_check": None,
        "error": None,
    }

    try:
        # curl -sL: follow redirects, silent mode
        proc = subprocess.run(
            ["curl", "-sL", "-o", "/dev/null", "-w",
             "%{http_code}|%{url_effective}|%{redirect_url}",
             "--max-time", str(timeout), url],
            capture_output=True, text=True, timeout=timeout + 5
        )
        parts = proc.stdout.strip().split("|")
        if len(parts) >= 2:
            result["status_code"] = int(parts[0]) if parts[0].isdigit() else None
            result["final_url"] = parts[1] if parts[1] else url
            result["redirect"] = result["final_url"] != url
            result["ok"] = result["status_code"] in (200, 301, 302)

        # 컨텐츠 키워드 매칭: 실제 HTML 가져와서 확인
        if result["ok"]:
            proc2 = subprocess.run(
                ["curl", "-sL", "--max-time", str(timeout), url],
                capture_output=True, text=True, timeout=timeout + 5
            )
            body = proc2.stdout[:10000].lower()  # 앞부분만 확인

            has_property = any(kw in body for kw in PROPERTY_KEYWORDS)
            has_redirect = any(kw in body for kw in REDIRECT_KEYWORDS)

            if has_redirect and not has_property:
                result["content_check"] = "WARN_REDIRECT_ROOT"
                result["ok"] = False
            elif has_property:
                result["content_check"] = "OK_PROPERTY_CONTENT"
            else:
                result["content_check"] = "OK_GENERIC"

    except subprocess.TimeoutExpired:
        result["error"] = "timeout"
    except Exception as e:
        result["error"] = str(e)

    return result


def validate_all_links(links: list, delay: float = 0.5) -> list:
    """모든 링크를 순차적으로 검증 (rate limiting)"""
    results = []
    seen = set()
    for i, link in enumerate(links):
        url = link["url"]
        if url in seen:
            continue
        seen.add(url)

        print(f"  [{i+1}/{len(links)}] {url[:80]}...", end=" ", flush=True)
        check = check_url(url)
        check["link_text"] = link.get("text", "")
        check["source"] = link.get("source", "")
        results.append(check)
        status = "OK" if check["ok"] else f"FAIL({check.get('status_code', check.get('error', '?'))})"
        if check.get("content_check") == "WARN_REDIRECT_ROOT":
            status += " [매물->검색루트]"
        print(status)

        time.sleep(delay)

    return results


def main():
    parser = argparse.ArgumentParser(description="외부 URL 전수 검증 스크립트")
    parser.add_argument("--docx", default=None, help="검증할 docx 파일")
    parser.add_argument("--html", default=None, help="검증할 HTML 파일")
    parser.add_argument("--output", default=None, help="검증 리포트 JSON 출력 경로")
    parser.add_argument("--delay", type=float, default=0.5, help="요청 간 딜레이 (초, 기본 0.5)")
    parser.add_argument("--timeout", type=int, default=15, help="URL 요청 타임아웃 (초, 기본 15)")
    args = parser.parse_args()

    if not args.docx and not args.html:
        print("오류: --docx 또는 --html 중 하나 이상 지정해야 합니다.", file=sys.stderr)
        sys.exit(1)

    all_links = []

    if args.docx:
        docx_path = Path(args.docx)
        if not docx_path.exists():
            print(f"오류: 파일을 찾을 수 없습니다: {docx_path}", file=sys.stderr)
            sys.exit(1)
        print(f"DOCX 하이퍼링크 추출 중: {docx_path}")
        docx_links = extract_links_from_docx(str(docx_path))
        print(f"  -> {len(docx_links)}개 링크 발견")
        all_links.extend(docx_links)

    if args.html:
        html_path = Path(args.html)
        if not html_path.exists():
            print(f"오류: 파일을 찾을 수 없습니다: {html_path}", file=sys.stderr)
            sys.exit(1)
        print(f"HTML 하이퍼링크 추출 중: {html_path}")
        html_links = extract_links_from_html(str(html_path))
        print(f"  -> {len(html_links)}개 링크 발견")
        all_links.extend(html_links)

    if not all_links:
        print("링크가 없습니다.")
        sys.exit(0)

    # 중복 제거 후 URL 검증
    unique_urls = list({link["url"]: link for link in all_links}.values())
    print(f"\n총 {len(unique_urls)}개 고유 URL 검증 시작...")
    results = validate_all_links(unique_urls, delay=args.delay)

    # 요약
    ok_count = sum(1 for r in results if r["ok"])
    fail_count = len(results) - ok_count
    redirect_warns = sum(1 for r in results if r.get("content_check") == "WARN_REDIRECT_ROOT")

    report = {
        "summary": {
            "total_links": len(all_links),
            "unique_urls": len(unique_urls),
            "ok": ok_count,
            "fail": fail_count,
            "redirect_root_warnings": redirect_warns,
        },
        "results": results,
        "failed": [r for r in results if not r["ok"]],
    }

    print(f"\n=== 검증 완료 ===")
    print(f"  전체: {len(unique_urls)} | 정상: {ok_count} | 실패: {fail_count} | 리다이렉트경고: {redirect_warns}")

    if args.output:
        out_path = Path(args.output)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(report, f, ensure_ascii=False, indent=2)
        print(f"리포트 저장: {out_path}")
    else:
        if report["failed"]:
            print("\n=== 실패 URL 목록 ===")
            for r in report["failed"]:
                print(f"  {r['status_code'] or r['error']} | {r['url'][:80]}")


if __name__ == "__main__":
    main()
